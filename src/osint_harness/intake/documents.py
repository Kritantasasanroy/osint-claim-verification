import email
import email.policy
import io
from abc import ABC, abstractmethod
from dataclasses import dataclass
from enum import StrEnum
from html.parser import HTMLParser
from pathlib import Path
from typing import ClassVar

import openpyxl
import pptx
import pypdf
from docx.package import Package as WordPackage
from PIL import Image, ImageSequence, UnidentifiedImageError
from pptx.enum.shapes import MSO_SHAPE_TYPE

from osint_harness.intake.ocr import NemotronParseOcr


class UnsupportedDocument(ValueError):
    def __init__(self, path: Path, supported: list[str]) -> None:
        super().__init__(
            f"Cannot read '{path.name}': the '{path.suffix or 'no extension'}' format is not supported. "
            f"Supported: {', '.join(supported)}. Convert legacy .doc/.xls/.ppt files to their modern format."
        )


class EmptyDocument(ValueError):
    def __init__(self, path: Path) -> None:
        super().__init__(f"No text could be recovered from '{path.name}', even with OCR.")


class ExtractionMethod(StrEnum):
    NATIVE_TEXT = "native_text"
    TEXT_LAYER = "text_layer"
    OCR = "ocr"
    UNREADABLE = "unreadable"


@dataclass(frozen=True)
class Passage:
    location: str
    text: str
    method: ExtractionMethod


@dataclass(frozen=True)
class Document:
    name: str
    passages: tuple[Passage, ...]

    @property
    def text(self) -> str:
        return "\n\n".join(f"[{passage.location}]\n{passage.text}" for passage in self.passages)

    @property
    def methods(self) -> frozenset[ExtractionMethod]:
        return frozenset(passage.method for passage in self.passages)

    def pieces(self, max_chars: int) -> list[Passage]:
        """The readable text cut into pieces of at most `max_chars`, each labelled with its origin."""
        return [
            Passage(
                f"{passage.location} part {start // max_chars + 1}",
                passage.text[start : start + max_chars],
                passage.method,
            )
            for passage in self.passages
            if passage.method is not ExtractionMethod.UNREADABLE
            for start in range(0, len(passage.text), max_chars)
        ]


class DocumentFormat(ABC):
    suffixes: ClassVar[frozenset[str]]
    MIN_EMBEDDED_IMAGE_SIDE: ClassVar[int] = 120

    @abstractmethod
    async def passages(self, path: Path, ocr: NemotronParseOcr) -> list[Passage]: ...

    async def _ocr_embedded(self, blob: bytes, location: str, ocr: NemotronParseOcr) -> list[Passage]:
        """An embedded picture in a format the imaging library can't decode (SVG, EMF) is noted, not fatal."""
        try:
            image = Image.open(io.BytesIO(blob))
            image.load()
        except (UnidentifiedImageError, OSError):
            note = "(embedded image in a format that could not be decoded; its content was not read)"
            return [Passage(location, note, ExtractionMethod.UNREADABLE)]
        if min(image.size) < self.MIN_EMBEDDED_IMAGE_SIDE:
            return []
        return [Passage(location, await ocr.transcribe(image), ExtractionMethod.OCR)]


class PlainTextFormat(DocumentFormat):
    suffixes = frozenset({".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".log", ".yaml", ".yml"})

    async def passages(self, path: Path, ocr: NemotronParseOcr) -> list[Passage]:
        return [Passage("text", self.decode(path.read_bytes()), ExtractionMethod.NATIVE_TEXT)]

    @staticmethod
    def decode(data: bytes) -> str:
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError:
            return data.decode("cp1252", errors="replace")


class _VisibleText(HTMLParser):
    HIDDEN: ClassVar[frozenset[str]] = frozenset({"script", "style", "noscript", "template"})

    def __init__(self) -> None:
        super().__init__()
        self.chunks: list[str] = []
        self._hidden_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self.HIDDEN:
            self._hidden_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self.HIDDEN and self._hidden_depth:
            self._hidden_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self._hidden_depth and data.strip():
            self.chunks.append(data.strip())


class HtmlFormat(DocumentFormat):
    suffixes = frozenset({".html", ".htm"})

    async def passages(self, path: Path, ocr: NemotronParseOcr) -> list[Passage]:
        parser = _VisibleText()
        parser.feed(PlainTextFormat.decode(path.read_bytes()))
        return [Passage("html", "\n".join(parser.chunks), ExtractionMethod.NATIVE_TEXT)]


class EmailFormat(DocumentFormat):
    suffixes = frozenset({".eml"})

    async def passages(self, path: Path, ocr: NemotronParseOcr) -> list[Passage]:
        message = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
        header = "\n".join(f"{name}: {message[name]}" for name in ("From", "To", "Date", "Subject"))
        body_part = message.get_body(preferencelist=("plain", "html"))
        body = body_part.get_content() if body_part is not None else ""
        return [Passage("email", f"{header}\n\n{body}", ExtractionMethod.NATIVE_TEXT)]


class PdfFormat(DocumentFormat):
    """Digital pages keep their exact text layer; scanned pages have their original images OCR'd."""

    suffixes = frozenset({".pdf"})
    MIN_TEXT_LAYER_CHARS: ClassVar[int] = 100

    async def passages(self, path: Path, ocr: NemotronParseOcr) -> list[Passage]:
        result: list[Passage] = []
        for number, page in enumerate(pypdf.PdfReader(path).pages, start=1):
            layout = page.extract_text(extraction_mode="layout")
            text = "\n".join(line.rstrip() for line in layout.splitlines()).strip()
            if text:
                result.append(Passage(f"page {number}", text, ExtractionMethod.TEXT_LAYER))
            if len(text) >= self.MIN_TEXT_LAYER_CHARS:
                continue
            for index, picture in enumerate(page.images, start=1):
                result.extend(await self._ocr_embedded(picture.data, f"page {number} scan {index}", ocr))
        return result


class ImageFormat(DocumentFormat):
    suffixes = frozenset({".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"})
    MAX_FRAMES: ClassVar[int] = 30

    async def passages(self, path: Path, ocr: NemotronParseOcr) -> list[Passage]:
        result: list[Passage] = []
        with Image.open(path) as image:
            for index, frame in enumerate(ImageSequence.Iterator(image), start=1):
                if index > self.MAX_FRAMES:
                    note = f"(frames after {self.MAX_FRAMES} were not read)"
                    result.append(Passage(f"image {index}+", note, ExtractionMethod.UNREADABLE))
                    break
                text = await ocr.transcribe(frame.copy())
                result.append(Passage(f"image {index}", text, ExtractionMethod.OCR))
        return result


class WordFormat(DocumentFormat):
    suffixes = frozenset({".docx"})

    async def passages(self, path: Path, ocr: NemotronParseOcr) -> list[Passage]:
        document = WordPackage.open(str(path)).main_document_part.document
        paragraphs = [str(paragraph.text) for paragraph in document.paragraphs if paragraph.text.strip()]
        tables = [
            " | ".join(str(cell.text).strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        result = [Passage("body", "\n".join(paragraphs + tables), ExtractionMethod.NATIVE_TEXT)]
        images = [rel.target_part.blob for rel in document.part.rels.values() if "image" in rel.reltype]
        for index, blob in enumerate(images, start=1):
            result.extend(await self._ocr_embedded(bytes(blob), f"embedded image {index}", ocr))
        return result


class SpreadsheetFormat(DocumentFormat):
    suffixes = frozenset({".xlsx", ".xlsm"})
    MAX_ROWS: ClassVar[int] = 2000

    async def passages(self, path: Path, ocr: NemotronParseOcr) -> list[Passage]:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        result: list[Passage] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(max_row=self.MAX_ROWS + 1, values_only=True))
            kept = rows[: self.MAX_ROWS]
            lines = [" | ".join(str(cell) for cell in row if cell is not None) for row in kept]
            text = "\n".join(line for line in lines if line)
            result.append(Passage(f"sheet {sheet.title}", text, ExtractionMethod.NATIVE_TEXT))
            if len(rows) > self.MAX_ROWS:
                note = f"(rows after {self.MAX_ROWS} were not read)"
                location = f"sheet {sheet.title} rows {self.MAX_ROWS + 1}+"
                result.append(Passage(location, note, ExtractionMethod.UNREADABLE))
        workbook.close()
        return result


class SlideFormat(DocumentFormat):
    suffixes = frozenset({".pptx"})

    async def passages(self, path: Path, ocr: NemotronParseOcr) -> list[Passage]:
        deck = pptx.Presentation(str(path))
        result: list[Passage] = []
        for number, slide in enumerate(deck.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(str(shape.text_frame.text))
                if shape.has_table:
                    texts.extend(" | ".join(str(cell.text) for cell in row.cells) for row in shape.table.rows)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    location = f"slide {number} image"
                    result.extend(await self._ocr_embedded(bytes(shape.image.blob), location, ocr))
            result.append(Passage(f"slide {number}", "\n".join(texts), ExtractionMethod.NATIVE_TEXT))
        return result


class DocumentIntake:
    FORMATS: ClassVar[tuple[DocumentFormat, ...]] = (
        PlainTextFormat(),
        HtmlFormat(),
        EmailFormat(),
        PdfFormat(),
        ImageFormat(),
        WordFormat(),
        SpreadsheetFormat(),
        SlideFormat(),
    )

    def __init__(self, ocr: NemotronParseOcr) -> None:
        self._ocr = ocr
        self._by_suffix = {suffix: fmt for fmt in self.FORMATS for suffix in fmt.suffixes}

    async def read(self, path: Path) -> Document:
        reader = self._by_suffix.get(path.suffix.lower())
        if reader is None:
            raise UnsupportedDocument(path, sorted(self._by_suffix))
        passages = tuple(p for p in await reader.passages(path, self._ocr) if p.text.strip())
        if all(passage.method is ExtractionMethod.UNREADABLE for passage in passages):
            raise EmptyDocument(path)
        return Document(name=path.name, passages=passages)
