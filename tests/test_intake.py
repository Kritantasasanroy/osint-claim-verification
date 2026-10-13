import asyncio
import json
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import ClassVar

import docx
import httpx
import openpyxl
import pptx
import pytest
from matplotlib.figure import Figure
from PIL import Image
from pydantic import ValidationError

from fakes import Handler
from osint_harness.intake.documents import (
    Document,
    DocumentIntake,
    EmptyDocument,
    ExtractionMethod,
    ImageFormat,
    SpreadsheetFormat,
    UnsupportedDocument,
)
from osint_harness.intake.ocr import NemotronParseOcr
from osint_harness.settings import Settings
from osint_harness.transport import ResponseCache, Transport


class OcrService:
    DIGITAL_TEXT: ClassVar[str] = (
        "Wirecard AG filed for insolvency in June 2020 after EUR 1.9 billion went missing from its accounts."
    )

    def __init__(self) -> None:
        self.calls = 0

    def handle(self, request: httpx.Request) -> httpx.Response:
        self.calls += 1
        arguments = json.dumps([{"text": f"OCR TEXT {self.calls}"}])
        return httpx.Response(
            200, json={"choices": [{"message": {"tool_calls": [{"function": {"arguments": arguments}}]}}]}
        )


@pytest.fixture
def ocr_service() -> OcrService:
    return OcrService()


@pytest.fixture
def intake(
    settings: Settings, make_transport: Callable[[Handler], Transport], ocr_service: OcrService
) -> DocumentIntake:
    return DocumentIntake(NemotronParseOcr(make_transport(ocr_service.handle), settings))


@pytest.fixture
def read(intake: DocumentIntake) -> Callable[[Path], Document]:
    return lambda path: asyncio.run(intake.read(path))


def test_digital_pdf_uses_its_text_layer_without_ocr(
    read: Callable[[Path], Document], ocr_service: OcrService, tmp_path: Path
) -> None:
    path = tmp_path / "filing.pdf"
    figure = Figure(figsize=(8.5, 11))
    figure.text(0.05, 0.9, OcrService.DIGITAL_TEXT, fontsize=8)
    figure.text(0.05, 0.8, "Its chief executive Markus Braun was arrested.", fontsize=8)
    figure.savefig(path)

    document = read(path)

    assert OcrService.DIGITAL_TEXT in document.text
    assert document.methods == {ExtractionMethod.TEXT_LAYER}
    assert ocr_service.calls == 0


def test_scanned_pdf_page_is_ocred(
    read: Callable[[Path], Document], ocr_service: OcrService, tmp_path: Path
) -> None:
    path = tmp_path / "scan.pdf"
    Image.new("RGB", (800, 1000), "white").save(path)

    document = read(path)

    assert "OCR TEXT 1" in document.text
    assert document.methods == {ExtractionMethod.OCR}


def test_tall_screenshot_is_tiled_for_ocr(
    read: Callable[[Path], Document], ocr_service: OcrService, tmp_path: Path
) -> None:
    path = tmp_path / "thread.png"
    Image.new("RGB", (600, 3000), "white").save(path)

    document = read(path)

    assert ocr_service.calls > 1
    assert "OCR TEXT 1" in document.text and f"OCR TEXT {ocr_service.calls}" in document.text


def test_word_document_reads_paragraphs_and_tables(read: Callable[[Path], Document], tmp_path: Path) -> None:
    path = tmp_path / "memo.docx"
    word = docx.Document()
    word.add_paragraph("Markus Braun was CEO of Wirecard.")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Auditor"
    table.rows[0].cells[1].text = "EY"
    word.save(str(path))

    text = read(path).text

    assert "Markus Braun was CEO of Wirecard." in text
    assert "Auditor | EY" in text


def test_spreadsheet_and_slides_are_read(read: Callable[[Path], Document], tmp_path: Path) -> None:
    sheet_path = tmp_path / "ledger.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Entity", "Cash"])
    workbook.active.append(["Wirecard", 1900000000])
    workbook.save(sheet_path)
    deck_path = tmp_path / "deck.pptx"
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "FTX collapse timeline"
    deck.save(str(deck_path))

    assert "Wirecard | 1900000000" in read(sheet_path).text
    assert "FTX collapse timeline" in read(deck_path).text


def test_html_drops_scripts(read: Callable[[Path], Document], tmp_path: Path) -> None:
    path = tmp_path / "page.html"
    path.write_text("<html><script>var x=1;</script><body><p>Theranos dissolved in 2018.</p></body></html>")

    text = read(path).text

    assert "Theranos dissolved in 2018." in text
    assert "var x" not in text


def test_unsupported_and_empty_files_are_reported(read: Callable[[Path], Document], tmp_path: Path) -> None:
    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"\xd0\xcf")
    blank = tmp_path / "blank.txt"
    blank.write_text("   ")

    with pytest.raises(UnsupportedDocument, match=r"\.doc"):
        read(legacy)
    with pytest.raises(EmptyDocument):
        read(blank)


def test_an_undecodable_embedded_image_is_noted_and_the_text_survives(
    read: Callable[[Path], Document], tmp_path: Path
) -> None:
    picture = tmp_path / "logo.png"
    Image.new("RGB", (300, 300), "white").save(picture)
    original = tmp_path / "memo.docx"
    word = docx.Document()
    word.add_paragraph("Board memo: Wirecard cash balances could not be confirmed.")
    word.add_picture(str(picture))
    word.save(str(original))
    with_svg = tmp_path / "memo_svg.docx"
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(with_svg, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("word/media/"):
                data = b'<svg xmlns="http://www.w3.org/2000/svg"><circle r="4"/></svg>'
            target.writestr(item, data)

    document = read(with_svg)

    assert "Wirecard cash balances could not be confirmed." in document.text
    assert ExtractionMethod.UNREADABLE in document.methods


def test_a_file_whose_only_content_is_undecodable_is_empty(
    read: Callable[[Path], Document], tmp_path: Path
) -> None:
    path = tmp_path / "logo.png"
    path.write_bytes(b'<svg xmlns="http://www.w3.org/2000/svg"><circle r="4"/></svg>')
    deck = tmp_path / "logo_only.pptx"
    presentation = pptx.Presentation()
    blank = presentation.slides.add_slide(presentation.slide_layouts[6])
    Image.new("RGB", (300, 300), "white").save(tmp_path / "real.png")
    blank.shapes.add_picture(str(tmp_path / "real.png"), 0, 0)
    presentation.save(str(deck))
    with_svg = tmp_path / "svg_only.pptx"
    with zipfile.ZipFile(deck) as source, zipfile.ZipFile(with_svg, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            target.writestr(item, path.read_bytes() if item.filename.startswith("ppt/media/") else data)

    with pytest.raises(EmptyDocument):
        read(with_svg)


def test_a_long_document_is_cut_into_readable_pieces_without_losing_text(
    read: Callable[[Path], Document], tmp_path: Path
) -> None:
    path = tmp_path / "filing.txt"
    path.write_text("x" * 9000)

    pieces = read(path).pieces(4000)

    assert [len(p.text) for p in pieces] == [4000, 4000, 1000]
    assert [p.location for p in pieces] == ["text part 1", "text part 2", "text part 3"]


def test_an_ocr_reply_without_text_blocks_is_not_replayed_from_cache(
    settings: Settings, tmp_path: Path
) -> None:
    calls: list[int] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(1)
        arguments = "not json" if len(calls) == 1 else json.dumps([{"text": "RECOVERED"}])
        return httpx.Response(
            200, json={"choices": [{"message": {"tool_calls": [{"function": {"arguments": arguments}}]}}]}
        )

    cache = ResponseCache(tmp_path / "c.sqlite3")
    ocr = NemotronParseOcr(
        Transport(httpx.AsyncClient(transport=httpx.MockTransport(handler)), cache, 1, 0.0), settings
    )
    image = Image.new("RGB", (400, 400), "white")

    with pytest.raises(ValidationError):
        asyncio.run(ocr.transcribe(image))
    assert asyncio.run(ocr.transcribe(image)) == "RECOVERED"


def test_a_spreadsheet_past_the_row_cap_says_what_was_not_read(
    read: Callable[[Path], Document], tmp_path: Path
) -> None:
    path = tmp_path / "big.xlsx"
    workbook = openpyxl.Workbook()
    for row in range(SpreadsheetFormat.MAX_ROWS + 5):
        workbook.active.append([f"row {row}"])
    workbook.save(path)

    document = read(path)

    assert ExtractionMethod.UNREADABLE in document.methods
    assert f"rows after {SpreadsheetFormat.MAX_ROWS} were not read" in document.text


def test_an_ocr_reply_without_a_tool_call_is_reported_and_not_replayed(
    settings: Settings, tmp_path: Path
) -> None:
    calls: list[int] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(1)
        if len(calls) == 1:
            return httpx.Response(200, json={"choices": [{"message": {"tool_calls": []}}]})
        arguments = json.dumps([{"text": "RECOVERED"}])
        return httpx.Response(
            200, json={"choices": [{"message": {"tool_calls": [{"function": {"arguments": arguments}}]}}]}
        )

    cache = ResponseCache(tmp_path / "c.sqlite3")
    ocr = NemotronParseOcr(
        Transport(httpx.AsyncClient(transport=httpx.MockTransport(handler)), cache, 1, 0.0), settings
    )
    image = Image.new("RGB", (400, 400), "white")

    with pytest.raises(ValidationError):
        asyncio.run(ocr.transcribe(image))
    assert asyncio.run(ocr.transcribe(image)) == "RECOVERED"


def test_an_image_past_the_frame_cap_says_what_was_not_read(
    read: Callable[[Path], Document], tmp_path: Path
) -> None:
    path = tmp_path / "scan.tif"
    frames = [Image.new("RGB", (200, 200), "white") for _ in range(ImageFormat.MAX_FRAMES + 2)]
    frames[0].save(path, save_all=True, append_images=frames[1:])

    document = read(path)

    assert f"frames after {ImageFormat.MAX_FRAMES} were not read" in document.text
